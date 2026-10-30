#!/usr/bin/env python3
"""Excel書き込みユーティリティ。JSON入力からExcelファイルを生成する。"""
import argparse
import json
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("ERROR: openpyxl が未インストールです。`pip install openpyxl` を実行してください。", file=sys.stderr)
    sys.exit(1)


def write_excel(out_path: str, sheet_name: str, data: dict, template: str = None) -> None:
    if template:
        wb = openpyxl.load_workbook(template)
        ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.active
        ws.title = sheet_name
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_name

    headers = data.get("headers", [])
    rows = data.get("rows", [])

    if headers:
        ws.append(headers)
    for row in rows:
        ws.append(row)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    print(f"書き込み完了: {out_path} (シート: {sheet_name}, {len(rows)} 行)")


def main():
    parser = argparse.ArgumentParser(description="JSONデータをExcelファイルに書き込む")
    parser.add_argument("--out", required=True, help="出力Excelファイルパス (.xlsx)")
    parser.add_argument("--sheet", default="Sheet1", help="書き込むシート名（デフォルト: Sheet1）")
    parser.add_argument("--data", help="データJSONファイルパス（省略時はstdinから読む）")
    parser.add_argument("--template", help="ベースにする既存Excelファイル (.xlsx)")
    args = parser.parse_args()

    if args.data:
        with open(args.data, encoding="utf-8") as f:
            data = json.load(f)
    else:
        data = json.load(sys.stdin)

    write_excel(args.out, args.sheet, data, args.template)


if __name__ == "__main__":
    main()
