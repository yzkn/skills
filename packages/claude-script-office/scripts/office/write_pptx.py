#!/usr/bin/env python3
"""PowerPoint書き込みユーティリティ。JSON入力からpptxファイルを生成する。"""
import argparse
import json
import sys
from pathlib import Path

try:
    from pptx import Presentation
except ImportError:
    print("ERROR: python-pptx が未インストールです。`pip install python-pptx` を実行してください。", file=sys.stderr)
    sys.exit(1)


def write_pptx(out_path: str, slides_data: dict, template: str = None) -> None:
    prs = Presentation(template) if template else Presentation()

    title_layout = prs.slide_layouts[0]
    content_layout = prs.slide_layouts[1]

    deck_title = slides_data.get("title")
    if deck_title:
        slide = prs.slides.add_slide(title_layout)
        slide.shapes.title.text = deck_title
        if len(slide.placeholders) > 1:
            slide.placeholders[1].text = slides_data.get("subtitle", "")

    for slide_data in slides_data.get("slides", []):
        slide = prs.slides.add_slide(content_layout)
        slide.shapes.title.text = slide_data.get("title", "")
        body_text = slide_data.get("body", "")
        if body_text and len(slide.placeholders) > 1:
            tf = slide.placeholders[1].text_frame
            tf.text = ""
            for line in body_text.split("\n"):
                if line.strip():
                    p = tf.add_paragraph()
                    p.text = line.strip()

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out_path)
    slide_count = len(slides_data.get("slides", []))
    print(f"書き込み完了: {out_path} ({slide_count} スライド)")


def main():
    parser = argparse.ArgumentParser(description="JSONデータからPowerPointファイルを生成する")
    parser.add_argument("--out", required=True, help="出力PowerPointファイルパス (.pptx)")
    parser.add_argument("--slides", help="スライドJSONファイルパス（省略時はstdinから読む）")
    parser.add_argument("--template", help="ベースにする既存PowerPointファイル (.pptx)")
    args = parser.parse_args()

    if args.slides:
        with open(args.slides, encoding="utf-8") as f:
            slides_data = json.load(f)
    else:
        slides_data = json.load(sys.stdin)

    write_pptx(args.out, slides_data, args.template)


if __name__ == "__main__":
    main()
