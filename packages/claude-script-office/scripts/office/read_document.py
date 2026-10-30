#!/usr/bin/env python3
"""読み取りディスパッチャ: anydoc を一次変換とし、使えない・品質が低い場合に
markitdown へ、それでも品質が低い場合に Docling へ自動フォールバックする。

エンジンの序列（自動判定時）:
    1. anydoc（@firecrawl/anydoc・Node CLI）… 既定の一次変換
    2. markitdown（Python）………………………… anydoc が未導入/未対応/失敗/品質 NG のとき
    3. Docling（Python）……………………………… 1・2 のいずれも品質 NG のとき（OCR もここが担う）
    anytomd（Rust CLI）は自動判定に入れず、--force-engine とクロスチェックの検証にだけ使う。

使い方:
    python read_document.py <ファイルパス>
    python read_document.py <ファイルパス> --force-engine anydoc|markitdown|docling|anytomd

出力:
    Markdown を stdout に出力する。先頭に使用エンジンを示すコメントを付す。
    判定・警告は stderr に出力する。
"""

import argparse
import subprocess
import sys
import zipfile
from pathlib import Path

# --- 自動フォールバックの閾値（実測でチューニング可能） -----------------------
MIN_EXTRACTION_RATIO = 0.30   # D-3: 出力文字数 / ソース推定文字数 の下限
SHORT_LINE_LEN = 25           # D-4: 「短行」とみなす文字数
SHORT_LINE_STREAK = 8         # D-4: 短行が連続する本数の閾値

# --- anydoc（一次変換エンジン） ----------------------------------------------
# 無スコープの `anydoc` は同名の別パッケージ（2018 年の web server）なので使わない。
ANYDOC_PACKAGE = "@firecrawl/anydoc"
ANYDOC_INSTALL_HINT = f"npm install -g {ANYDOC_PACKAGE}"
ANYDOC_EXIT_NEEDS_OCR = 3   # `--ocr reject`（既定）で OCR が必要な PDF を拒否したときの終了コード
# anydoc が変換できる拡張子（README のサポート表）。これ以外は markitdown が担当する。
ANYDOC_EXTS = {
    ".doc", ".docx", ".docm",
    ".ppt", ".pps", ".pot", ".pptx", ".pptm", ".ppsx", ".ppsm",
    ".xls", ".xlsx", ".xlsm", ".xlsb",
    ".odt", ".ods", ".odp",
    ".rtf", ".epub", ".csv", ".pdf",
}


def _md_table_lines(md: str):
    """Markdown テーブル行（| で始まり | を含む行）を返す。"""
    lines = []
    for ln in md.splitlines():
        s = ln.strip()
        if s.startswith("|") and s.count("|") >= 2:
            lines.append(s)
    return lines


def _md_tables(md: str):
    """連続する表行を 1 つの表としてまとめ、表ごとの行リストを返す。

    複数の表（Excel のシートごとの表など）を 1 つにまとめて列数を比べると、
    表ごとに列数が違うだけで「列数が不揃い」と誤判定する。表単位に区切る。
    """
    tables, current = [], []
    for ln in md.splitlines():
        s = ln.strip()
        if s.startswith("|") and s.count("|") >= 2:
            current.append(s)
            continue
        if current:
            tables.append(current)
            current = []
    if current:
        tables.append(current)
    return tables


def source_has_table(path: Path) -> bool:
    """ソースが表を含むかを推定する。ライブラリがあれば厳密に、無ければ拡張子で保守的に判定。"""
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            import docx  # python-docx
            return len(docx.Document(str(path)).tables) > 0
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True)
            for ws in wb.worksheets:
                if (ws.max_row or 0) >= 2 and (ws.max_column or 0) >= 2:
                    return True
            return False
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(path))
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_table:
                        return True
            return False
    except Exception as e:
        print(f"INFO: 表の有無の厳密判定に失敗（拡張子で推定します）: {e}", file=sys.stderr)
    # ライブラリ不在・その他: xlsx/pptx/docx は表を含む可能性が高いと保守的にみなす
    return ext in (".xlsx", ".xlsm", ".pptx", ".docx")


def source_text_len(path: Path):
    """ソースのプレーンテキスト文字数を推定する。推定できない場合は None。"""
    ext = path.suffix.lower()
    try:
        if ext == ".docx":
            import docx
            d = docx.Document(str(path))
            n = sum(len(p.text) for p in d.paragraphs)
            for t in d.tables:
                for row in t.rows:
                    for cell in row.cells:
                        n += len(cell.text)
            return n
        if ext in (".xlsx", ".xlsm"):
            import openpyxl
            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            n = 0
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    for v in row:
                        if v is not None:
                            n += len(str(v))
            return n
        if ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(str(path))
            n = 0
            for slide in prs.slides:
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        n += len(shape.text_frame.text)
            return n
    except Exception:
        return None
    return None  # PDF 等は推定しない（D-3 をスキップ）


def check_quality(md: str, path: Path):
    """抽出品質を評価する。フォールバックすべき理由のリストを返す（空なら合格）。"""
    reasons = []
    table_lines = _md_table_lines(md)

    # D-1: ソースに表があるのに Markdown 表が 0 件
    if source_has_table(path) and len(table_lines) == 0:
        reasons.append("D-1: ソースに表があるが Markdown 表が抽出されていない")

    # D-2: Markdown 表の各行のセル数が不揃い（表ごとに評価する。表をまたぐ列数の違いは正常）
    for i, tbl in enumerate(_md_tables(md), 1):
        # 区切り行(---)を除外して評価
        content_counts = [ln.count("|") for ln in tbl if set(ln.replace("|", "").strip()) - set("-: ")]
        if len(set(content_counts)) > 1:
            reasons.append(f"D-2: {i} 番目の Markdown 表で列数が行間で不揃い（結合セル・改行崩れの兆候）")
            break

    # D-3: 抽出文字数がソース推定の一定割合を下回る
    src_len = source_text_len(path)
    if src_len and src_len > 0:
        out_len = len("".join(md.split()))
        ratio = out_len / src_len
        if ratio < MIN_EXTRACTION_RATIO:
            reasons.append(f"D-3: 抽出文字数が少ない（{ratio:.0%} < {MIN_EXTRACTION_RATIO:.0%}）")

    # D-4: PDF の多段組・読み順崩れ（短行の連続を保守的に検出）
    if path.suffix.lower() == ".pdf":
        streak = max_streak = 0
        for ln in md.splitlines():
            s = ln.strip()
            if 0 < len(s) <= SHORT_LINE_LEN:
                streak += 1
                max_streak = max(max_streak, streak)
            else:
                streak = 0
        if max_streak >= SHORT_LINE_STREAK:
            reasons.append(f"D-4: 短行が {max_streak} 連続（多段組の読み順崩れの兆候）")

    return reasons


def _anydoc_command():
    """anydoc CLI の実行コマンドを返す。未導入なら None。

    グローバルインストール（`npm install -g @firecrawl/anydoc`）を前提とし、npx での
    都度取得はしない（実行のたびにレジストリへ取りに行くため）。
    """
    from shutil import which
    exe = which("anydoc")
    return [exe] if exe else None


def run_anydoc(path: Path):
    """anydoc を実行して (Markdown, ステータス) を返す。

    ステータス:
        ok         … 変換成功
        missing    … anydoc が未導入（グローバルインストールが必要）
        needs_ocr  … スキャン PDF 等で OCR が必要（anydoc はローカル OCR を持たない）
        error      … 上記以外の失敗（未対応形式・暗号化・破損・実行時エラー）

    OCR が必要な PDF は既定（`--ocr reject`）で終了コード 3 になる。`--ocr hosted` は
    文書を Firecrawl Parse へ送信するため、本スクリプトからは決して指定しない
    （外部送信の禁止。`.claude/rules/external-resources.md`）。
    """
    cmd = _anydoc_command()
    if cmd is None:
        return None, "missing"
    try:
        r = subprocess.run([*cmd, str(path)], capture_output=True)
    except Exception as e:
        print(f"WARNING: anydoc の実行に失敗しました: {e}", file=sys.stderr)
        return None, "error"
    if r.returncode != 0:
        err = _decode_bytes(r.stderr).strip()
        # 終了コード 3 = OCR 必要。将来コードが変わっても拾えるよう文言でも判定する。
        if r.returncode == ANYDOC_EXIT_NEEDS_OCR or "needsocr" in err.lower().replace(" ", "").replace("_", ""):
            print(f"INFO: anydoc: OCR が必要な文書です（anydoc はローカル OCR を持ちません）: {err[:200]}",
                  file=sys.stderr)
            return None, "needs_ocr"
        print(f"WARNING: anydoc が異常終了しました（exit={r.returncode}）: {err[:200]}", file=sys.stderr)
        return None, "error"
    md = _decode_bytes(r.stdout)
    if not md.strip():
        # 正常終了でも中身が空なら採用しない（品質判定は空文字列を通してしまうため）
        print("WARNING: anydoc の出力が空でした。", file=sys.stderr)
        return None, "error"
    return md, "ok"


# --- anytomd（クロスチェック用の独立実装） ------------------------------------
# Rust 製（calamine / quick-xml）。openpyxl・python-pptx を使わないため、markitdown 一次の
# xlsx / pptx を独立に検証できる。自動フォールバックの順序には入れず、--force-engine と
# verify_conversion.py の検証エンジンとしてのみ使う。
ANYTOMD_INSTALL_HINT = "cargo install anytomd --locked --no-default-features"
# --gemini は画像を Gemini API へ送る（外部送信）。本スクリプトは渡さず、API キーも子プロセスへ渡さない。
ANYTOMD_SCRUB_ENV = ("GEMINI_API_KEY", "GOOGLE_API_KEY")


def _anytomd_command():
    """anytomd CLI の実行コマンドを返す。未導入なら None。

    `cargo install` の既定の導入先 `~/.cargo/bin` は PATH に載っていないことがあるため、
    PATH で見つからなければそこも探す。
    """
    from shutil import which
    exe = which("anytomd")
    if exe:
        return [exe]
    for name in ("anytomd.exe", "anytomd"):
        p = Path.home() / ".cargo" / "bin" / name
        if p.exists():
            return [str(p)]
    return None


def run_anytomd(path: Path):
    """anytomd を実行して Markdown を返す。失敗・未導入時は None。"""
    cmd = _anytomd_command()
    if cmd is None:
        print(f"ERROR: anytomd が未導入です（{ANYTOMD_INSTALL_HINT}）。", file=sys.stderr)
        return None
    import os
    env = {k: v for k, v in os.environ.items() if k not in ANYTOMD_SCRUB_ENV}
    try:
        r = subprocess.run([*cmd, str(path)], capture_output=True, env=env)
    except Exception as e:
        print(f"WARNING: anytomd の実行に失敗しました: {e}", file=sys.stderr)
        return None
    if r.returncode != 0:
        err = _decode_bytes(r.stderr).strip()[:200]
        print(f"WARNING: anytomd が異常終了しました（exit={r.returncode}）: {err}", file=sys.stderr)
        return None
    md = _decode_bytes(r.stdout)
    if not md.strip():
        print("WARNING: anytomd の出力が空でした。", file=sys.stderr)
        return None
    return md


def _markitdown_command():
    """利用可能な markitdown 実行コマンドを返す。無ければ None。

    Windows では CLI がシェルの PATH に載らないことがあるため、`python -m markitdown`
    （モジュール実行）を優先し、次に bare CLI を試す。
    """
    # 1) python -m markitdown（最も堅牢）
    try:
        import markitdown  # noqa: F401
        return [sys.executable, "-m", "markitdown"]
    except Exception:
        pass
    # 2) bare CLI
    from shutil import which
    if which("markitdown"):
        return ["markitdown"]
    return None


def run_markitdown(path: Path):
    """markitdown を実行して Markdown を返す。失敗・未導入時は None。"""
    cmd = _markitdown_command()
    if cmd is None:
        print("ERROR: markitdown が未インストールです（pip install markitdown）。", file=sys.stderr)
        return None
    import os
    # Windows では子プロセスの stdout が locale エンコーディング（cp932 等）になり
    # UTF-8 デコードに失敗するため、PYTHONIOENCODING=utf-8 を明示し、バイトで受けて堅牢に復号する。
    env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
    try:
        r = subprocess.run([*cmd, str(path)], capture_output=True, env=env)
    except Exception as e:
        print(f"WARNING: markitdown の実行に失敗しました: {e}", file=sys.stderr)
        return None
    if r.returncode != 0:
        err = _decode_bytes(r.stderr).strip()[:200]
        print(f"WARNING: markitdown が異常終了しました: {err}", file=sys.stderr)
        return None
    return _decode_bytes(r.stdout)


def _decode_bytes(b: bytes) -> str:
    """UTF-8 を優先し、失敗時は locale エンコーディング、最後は置換で復号する。"""
    if b is None:
        return ""
    for enc in ("utf-8", "cp932", "shift_jis"):
        try:
            return b.decode(enc)
        except Exception:
            continue
    return b.decode("utf-8", errors="replace")


def docling_available() -> bool:
    try:
        import docling  # noqa: F401
        return True
    except Exception:
        return False


def run_docling(path: Path):
    """Docling で変換して Markdown を返す。失敗時は None。"""
    try:
        from docling.document_converter import DocumentConverter
        conv = DocumentConverter()
        result = conv.convert(str(path))
        return result.document.export_to_markdown()
    except Exception as e:
        print(f"WARNING: Docling 変換に失敗しました: {e}", file=sys.stderr)
        return None


def _init_stdout():
    """stdout を UTF-8 / LF 固定にする。

    Windows の既定では print が \\n を \\r\\n に変換するため、変換結果に元からあった
    \\r\\n が \\r\\r\\n になり、リダイレクト保存したファイルで表の各行が空行に分断される。
    また既定エンコーディングが cp932 だと日本語がリダイレクト先で化ける。
    """
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", newline="\n")
        except Exception:
            pass  # 古い Python・差し替え済みストリームでは既定動作にフォールバック


def _print_reasons(engine: str, reasons: list):
    for r in reasons:
        print(f"INFO: 品質判定 NG（{engine}） → {r}", file=sys.stderr)


# anydoc を一次変換に使えなかった理由（run_anydoc のステータスに対応）
ANYDOC_SKIP_MSG = {
    "missing": f"anydoc が未導入です（{ANYDOC_INSTALL_HINT} でグローバルインストールしてください）",
    "needs_ocr": "OCR が必要な文書です（OCR は Docling が担当します）",
    "error": "anydoc で変換できませんでした",
}


def auto_convert(path: Path):
    """anydoc → markitdown → Docling の順に試し、(Markdown, engine ラベル) を返す。

    どのエンジンでも変換できなかった場合は (None, None) を返す。
    """
    ext = path.suffix.lower()
    anydoc_md = None

    if ext in ANYDOC_EXTS:
        md, status = run_anydoc(path)
        if status == "ok" and md is not None:
            reasons = check_quality(md, path)
            if not reasons:
                return md, "engine=anydoc"
            _print_reasons("anydoc", reasons)
            anydoc_md = md
        else:
            print(f"INFO: {ANYDOC_SKIP_MSG[status]}。", file=sys.stderr)
        print("INFO: markitdown へフォールバックします。", file=sys.stderr)
    else:
        print(f"INFO: {ext or '(拡張子なし)'} は anydoc の対応形式外のため markitdown を使用します。", file=sys.stderr)

    md = run_markitdown(path)
    if md is not None:
        reasons = check_quality(md, path)
        if not reasons:
            label = "engine=markitdown" if anydoc_md is None else "engine=markitdown (anydoc からフォールバック)"
            return md, label
        _print_reasons("markitdown", reasons)

    # anydoc・markitdown のいずれも品質 NG（または実行不可）→ Docling
    if docling_available():
        print("INFO: Docling へフォールバックします。", file=sys.stderr)
        dmd = run_docling(path)
        if dmd is not None:
            return dmd, "engine=docling (自動フォールバック)"
        print("WARNING: Docling 変換に失敗しました。", file=sys.stderr)
    else:
        print("WARNING: 抽出品質が低い可能性があります。Docling 導入で改善します（pip install docling）。",
              file=sys.stderr)

    # Docling が使えない場合は、得られた出力のうち精度の高い anydoc を優先して返す
    if anydoc_md is not None:
        print("WARNING: 品質判定 NG のまま anydoc の結果を使用します。", file=sys.stderr)
        return anydoc_md, "engine=anydoc (品質判定 NG)"
    if md is not None:
        print("WARNING: 品質判定 NG のまま markitdown の結果を使用します。", file=sys.stderr)
        return md, "engine=markitdown (品質判定 NG)"
    return None, None


def forced_convert(path: Path, engine: str):
    """--force-engine で指定されたエンジンだけを実行する（フォールバックしない）。"""
    if engine == "anydoc":
        md, status = run_anydoc(path)
        if status == "missing":
            print(f"ERROR: anydoc が未導入です。`{ANYDOC_INSTALL_HINT}` でグローバルインストールしてください。",
                  file=sys.stderr)
        return md
    if engine == "markitdown":
        return run_markitdown(path)
    if engine == "anytomd":
        return run_anytomd(path)
    return run_docling(path)


def main():
    _init_stdout()
    ap = argparse.ArgumentParser(
        description="Office/PDF を Markdown に変換する（anydoc→markitdown→Docling 自動フォールバック）")
    ap.add_argument("path", help="入力ファイルパス")
    ap.add_argument("--force-engine", choices=["anydoc", "markitdown", "docling", "anytomd"],
                    help="エンジンを固定する（自動判定・フォールバックを無効化）")
    args = ap.parse_args()

    path = Path(args.path)
    if not path.exists():
        print(f"ERROR: ファイルが存在しません: {path}", file=sys.stderr)
        sys.exit(2)

    # --- エンジン固定 ---
    if args.force_engine:
        md = forced_convert(path, args.force_engine)
        if md is None:
            sys.exit(1)
        print(f"<!-- read_document: engine={args.force_engine} (forced) -->")
        print(md)
        return

    # --- 自動判定 ---
    md, label = auto_convert(path)
    if md is None:
        sys.exit(1)
    print(f"<!-- read_document: {label} -->")
    print(md)


if __name__ == "__main__":
    main()
