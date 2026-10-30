#!/usr/bin/env python3
"""Word文書書き込みユーティリティ。JSON入力からdocxファイルを生成する。"""
import argparse
import json
import sys
from pathlib import Path

try:
    from docx import Document
except ImportError:
    print("ERROR: python-docx が未インストールです。`pip install python-docx` を実行してください。", file=sys.stderr)
    sys.exit(1)


def add_table(doc: Document, table_data: dict) -> None:
    headers = table_data.get("headers", [])
    rows = table_data.get("rows", [])
    col_count = len(headers) or (len(rows[0]) if rows else 0)
    if col_count == 0:
        return
    table = doc.add_table(rows=1 + len(rows), cols=col_count)
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.cell(0, i).text = str(header)
    for r_idx, row in enumerate(rows):
        for c_idx, cell_val in enumerate(row):
            table.cell(r_idx + 1, c_idx).text = str(cell_val)


def write_word(out_path: str, content: dict, template: str = None) -> None:
    doc = Document(template) if template else Document()

    title = content.get("title")
    if title:
        doc.add_heading(title, level=0)

    for section in content.get("sections", []):
        heading = section.get("heading")
        body = section.get("body")
        table_data = section.get("table")
        level = section.get("level", 1)

        if heading:
            doc.add_heading(heading, level=level)
        if body:
            doc.add_paragraph(body)
        if table_data:
            add_table(doc, table_data)

    Path(out_path).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    section_count = len(content.get("sections", []))
    print(f"書き込み完了: {out_path} ({section_count} セクション)")


def main():
    parser = argparse.ArgumentParser(description="JSONデータからWordファイルを生成する")
    parser.add_argument("--out", required=True, help="出力Wordファイルパス (.docx)")
    parser.add_argument("--content", help="コンテンツJSONファイルパス（省略時はstdinから読む）")
    parser.add_argument("--template", help="ベースにする既存Wordファイル (.docx)")
    args = parser.parse_args()

    if args.content:
        with open(args.content, encoding="utf-8") as f:
            content = json.load(f)
    else:
        content = json.load(sys.stdin)

    write_word(args.out, content, args.template)


if __name__ == "__main__":
    main()
