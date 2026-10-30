#!/usr/bin/env python3
"""デザインシステム対応 PowerPoint 生成ユーティリティ。

JSON 入力からコンサルティング品質の .pptx を生成する。
8 種類のレイアウト（title / section / content / two-column / table / chart / image / closing）をサポート。
画像埋め込みは Base64 data URI・URL・ファイルパスの 3 形式に対応する。
デザインは --theme でテーマ（slides-maker スキルの themes/<名前>.json）を切り替えられる（既定: royal-navy）。
"""

import argparse
import base64
import io
import json
import sys
import urllib.request
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.chart.data import ChartData
except ImportError:
    print("ERROR: python-pptx が未インストールです。`pip install python-pptx` を実行してください。", file=sys.stderr)
    sys.exit(1)


# ---------------------------------------------------------------------------
# Design System Constants
# ---------------------------------------------------------------------------

class Colors:
    NAVY_DARK        = RGBColor(0x1B, 0x3A, 0x6B)  # #1B3A6B ─ プライマリ
    NAVY_MID         = RGBColor(0x3D, 0x60, 0x98)  # #3D6098
    BLUE_LIGHT       = RGBColor(0x7A, 0x9C, 0xC6)  # #7A9CC6
    BLUE_PALE        = RGBColor(0xB8, 0xCC, 0xE4)  # #B8CCE4
    BLUE_ULTRA_PALE  = RGBColor(0xD9, 0xE5, 0xF2)  # #D9E5F2
    WHITE            = RGBColor(0xFF, 0xFF, 0xFF)  # #FFFFFF
    TEXT_DARK        = RGBColor(0x1A, 0x1A, 0x1A)  # #1A1A1A
    TEXT_GRAY        = RGBColor(0x76, 0x76, 0x76)  # #767676
    GRAY_LIGHT       = RGBColor(0xF5, 0xF7, 0xFA)  # #F5F7FA
    ACCENT_RED       = RGBColor(0xC0, 0x39, 0x2B)  # #C0392B ─ 警告・マイナス
    ACCENT_GREEN     = RGBColor(0x27, 0xAE, 0x60)  # #27AE60 ─ ポジティブ


class Fonts:
    HEADING = "游明朝"      # 見出し（セリフ体）
    BODY    = "游ゴシック"  # 本文（サンセリフ体）
    NUMBERS = "Calibri"    # 数値・英数字


# スライドサイズ: 16:9 ワイドスクリーン
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

# アクションタイトルバーの高さ
TITLE_BAR_H = Inches(1.1)

# チャート系列カラーパレット（ネイビー系グラデーション）
CHART_PALETTE = [
    Colors.NAVY_DARK,
    Colors.NAVY_MID,
    Colors.BLUE_LIGHT,
    Colors.BLUE_PALE,
    Colors.TEXT_GRAY,
]


# ---------------------------------------------------------------------------
# テーマ切り替え
# 上記の Colors / Fonts / CHART_PALETTE は既定テーマ（royal-navy）の値。
# テーマ定義は slides-maker スキル配下の themes/<名前>.json で管理し、
# --theme で切り替える（未指定キーは既定値のまま）。
# ---------------------------------------------------------------------------

THEMES_DIR = Path(__file__).resolve().parents[2] / "skills" / "slides-maker" / "themes"
DEFAULT_THEME = "royal-navy"


def load_theme(spec: str):
    """テーマ名（themes/<名前>.json）または JSON パスからテーマ定義を読み込む"""
    path = Path(spec) if spec.endswith(".json") else THEMES_DIR / f"{spec}.json"
    if not path.exists():
        if spec == DEFAULT_THEME:
            # 既定テーマファイルが無い環境ではスクリプト内蔵の既定値で続行する
            print(f"NOTE: 既定テーマ定義 {path} が無いため内蔵の既定値を使用します", file=sys.stderr)
            return None
        available = sorted(p.stem for p in THEMES_DIR.glob("*.json")) if THEMES_DIR.is_dir() else []
        print(f"ERROR: テーマが見つかりません: {spec}（利用可能: {available}）", file=sys.stderr)
        sys.exit(1)
    return json.loads(path.read_text(encoding="utf-8"))


def apply_theme(theme: dict) -> None:
    """テーマ定義で Colors / Fonts / CHART_PALETTE を上書きする"""
    global CHART_PALETTE
    for key, val in theme.get("colors", {}).items():
        setattr(Colors, key.upper(), RGBColor.from_string(val))
    fonts = theme.get("fonts", {}).get("python-pptx", {})
    Fonts.HEADING = fonts.get("heading", Fonts.HEADING)
    Fonts.BODY = fonts.get("body", Fonts.BODY)
    Fonts.NUMBERS = fonts.get("numbers", Fonts.NUMBERS)
    palette = theme.get("chart_palette", {}).get("python-pptx")
    if palette:
        # 色名（colors のキー）または HEX 直書きの両方を許容する
        CHART_PALETTE = [getattr(Colors, k.upper()) if hasattr(Colors, k.upper())
                         else RGBColor.from_string(k) for k in palette]


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _fill(shape, color: RGBColor) -> None:
    """シェイプを単色で塗りつぶす。"""
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _no_line(shape) -> None:
    """シェイプの枠線を非表示にする。"""
    shape.line.fill.background()


def _rect(slide, x, y, w, h, color: RGBColor):
    """塗りつぶし矩形を追加する。戻り値は shape。"""
    s = slide.shapes.add_shape(1, x, y, w, h)  # 1 = MSO_AUTO_SHAPE_TYPE.RECTANGLE
    _fill(s, color)
    _no_line(s)
    return s


def _textbox(slide, text: str, x, y, w, h,
             font=Fonts.BODY, size=18, bold=False,
             color=Colors.TEXT_DARK, align=PP_ALIGN.LEFT,
             wrap=True):
    """テキストボックスを追加する。戻り値は shape。"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return tb


def _action_title_bar(slide, title: str):
    """上部ネイビーバー＋アクションタイトルを追加する。"""
    _rect(slide, Inches(0), Inches(0), SLIDE_W, TITLE_BAR_H, Colors.NAVY_DARK)
    _textbox(
        slide, title,
        Inches(0.4), Inches(0.1), Inches(12.5), Inches(0.9),
        font=Fonts.HEADING, size=22, bold=True,
        color=Colors.WHITE, align=PP_ALIGN.LEFT,
    )


def _bullets(slide, items, x, y, w, h):
    """箇条書きテキストボックスを追加する。items は文字列またはdictのリスト。"""
    if not items:
        return
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        if isinstance(item, dict):
            text  = item.get("text", "")
            level = item.get("level", 0)
        else:
            text  = str(item)
            level = 0

        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = min(level, 4)
        run = p.add_run()
        run.text = text
        if level == 0:
            run.font.name = Fonts.BODY
            run.font.size = Pt(18)
            run.font.color.rgb = Colors.TEXT_DARK
            p.space_before = Pt(6)
            p.space_after  = Pt(3)
        else:
            run.font.name = Fonts.BODY
            run.font.size = Pt(15)
            run.font.color.rgb = Colors.TEXT_GRAY
            p.space_before = Pt(2)
            p.space_after  = Pt(2)


def _table_cell(cell, text: str, font=Fonts.BODY, size=13,
                bold=False, text_color=Colors.TEXT_DARK,
                bg_color=None, align=PP_ALIGN.LEFT):
    """テーブルセルにテキストとスタイルを設定する。"""
    if bg_color:
        cell.fill.solid()
        cell.fill.fore_color.rgb = bg_color

    tf = cell.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = text_color


def _load_image(src: str) -> io.BytesIO:
    """画像ソースから BytesIO を返す。

    対応形式:
      - data URI  : ``data:image/png;base64,<base64>``
      - URL       : ``http://`` / ``https://``
      - ファイルパス: それ以外の文字列
    """
    if src.startswith("data:"):
        # data:image/png;base64,XXXX
        _, encoded = src.split(",", 1)
        return io.BytesIO(base64.b64decode(encoded))
    elif src.startswith(("http://", "https://")):
        req = urllib.request.Request(src, headers={"User-Agent": "python-pptx-slides-maker/1.0"})
        with urllib.request.urlopen(req, timeout=15) as resp:
            return io.BytesIO(resp.read())
    else:
        with open(src, "rb") as f:
            return io.BytesIO(f.read())


def _add_image_fit(slide, src: str, max_left, max_top, max_w, max_h):
    """画像をアスペクト比を維持したまま指定矩形内にフィットして追加する。

    Returns:
        追加した Picture shape。読み込み失敗時は None。
    """
    try:
        stream = _load_image(src)
    except Exception as e:
        print(f"WARNING: 画像の読み込みに失敗しました ({src[:60]}): {e}", file=sys.stderr)
        return None

    # まず width = max_w で追加（高さは比例計算される）
    pic = slide.shapes.add_picture(stream, max_left, max_top, width=max_w)

    if pic.height > max_h:
        # 縦に収まらない場合は height 基準でリサイズ
        aspect = pic.width / pic.height
        pic.height = max_h
        pic.width  = int(max_h * aspect)
        # 水平中央揃え
        pic.left = max_left + (max_w - pic.width) // 2
    else:
        # 垂直中央揃え
        pic.top = max_top + (max_h - pic.height) // 2

    return pic


# ---------------------------------------------------------------------------
# Layout renderers
# ---------------------------------------------------------------------------

def render_title(slide, data: dict) -> None:
    """タイトルスライド: ダークネイビー背景、中央寄せタイトル。"""
    _rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, Colors.NAVY_DARK)
    # アクセントライン
    _rect(slide, Inches(0.8), Inches(3.5), Inches(4.0), Inches(0.05), Colors.BLUE_LIGHT)

    _textbox(slide, data.get("title", ""),
             Inches(0.8), Inches(2.0), Inches(11.0), Inches(1.4),
             font=Fonts.HEADING, size=36, bold=True,
             color=Colors.WHITE, align=PP_ALIGN.LEFT)

    subtitle = data.get("subtitle", "")
    if subtitle:
        _textbox(slide, subtitle,
                 Inches(0.8), Inches(3.7), Inches(10.5), Inches(0.8),
                 font=Fonts.BODY, size=20,
                 color=Colors.BLUE_PALE, align=PP_ALIGN.LEFT)

    meta = "  |  ".join(filter(None, [data.get("date", ""), data.get("author", "")]))
    if meta:
        _textbox(slide, meta,
                 Inches(0.8), Inches(6.5), Inches(10.5), Inches(0.6),
                 font=Fonts.BODY, size=13,
                 color=Colors.BLUE_LIGHT, align=PP_ALIGN.LEFT)


def render_section(slide, data: dict) -> None:
    """セクション区切り: 左ネイビーストライプ＋番号＋タイトル。"""
    _rect(slide, Inches(0), Inches(0), Inches(1.5), SLIDE_H, Colors.NAVY_DARK)

    section_num = str(data.get("section_number", ""))
    if section_num:
        _textbox(slide, section_num,
                 Inches(1.9), Inches(2.3), Inches(3.0), Inches(1.0),
                 font=Fonts.NUMBERS, size=52, bold=True,
                 color=Colors.BLUE_LIGHT, align=PP_ALIGN.LEFT)

    _textbox(slide, data.get("title", ""),
             Inches(1.9), Inches(3.2), Inches(10.5), Inches(1.5),
             font=Fonts.HEADING, size=30, bold=True,
             color=Colors.NAVY_DARK, align=PP_ALIGN.LEFT)

    subtitle = data.get("subtitle", "")
    if subtitle:
        _textbox(slide, subtitle,
                 Inches(1.9), Inches(4.8), Inches(10.5), Inches(0.8),
                 font=Fonts.BODY, size=18,
                 color=Colors.TEXT_GRAY, align=PP_ALIGN.LEFT)


def render_content(slide, data: dict) -> None:
    """コンテンツスライド: アクションタイトルバー＋箇条書き。

    ``image`` キーがある場合は左に箇条書き（60%幅）、右に画像（38%幅）の分割レイアウトになる。
    """
    title = data.get("action_title") or data.get("title", "")
    _action_title_bar(slide, title)

    image_data = data.get("image")
    src = ""
    if image_data:
        src = image_data if isinstance(image_data, str) else image_data.get("src", "")

    if src:
        # 左: 箇条書き (〜60%)、右: 画像 (〜38%)
        _bullets(slide, data.get("bullets", []),
                 Inches(0.4), Inches(1.25), Inches(7.6), Inches(5.9))
        _add_image_fit(slide, src,
                       Inches(8.2), Inches(1.3), Inches(4.9), Inches(5.9))
    else:
        _bullets(slide, data.get("bullets", []),
                 Inches(0.4), Inches(1.25), Inches(12.5), Inches(5.9))


def render_two_column(slide, data: dict) -> None:
    """2カラムスライド: アクションタイトルバー＋左右コンテンツ。"""
    title = data.get("action_title") or data.get("title", "")
    _action_title_bar(slide, title)

    # 仕切り線
    _rect(slide, Inches(6.65), Inches(1.3), Inches(0.03), Inches(5.9), Colors.BLUE_PALE)

    for side, x in (("left", Inches(0.4)), ("right", Inches(6.8))):
        col = data.get(side, {})
        if not col:
            continue
        heading = col.get("heading", "")
        if heading:
            _textbox(slide, heading, x, Inches(1.35), Inches(5.8), Inches(0.6),
                     font=Fonts.HEADING, size=16, bold=True,
                     color=Colors.NAVY_MID, align=PP_ALIGN.LEFT)
        _bullets(slide, col.get("bullets", []),
                 x, Inches(2.0), Inches(5.8), Inches(5.2))


def render_table(slide, data: dict) -> None:
    """テーブルスライド: ヘッダー行ネイビー＋交互ストライプ。"""
    title = data.get("action_title") or data.get("title", "")
    _action_title_bar(slide, title)

    headers = data.get("headers", [])
    rows    = data.get("rows", [])
    if not headers and not rows:
        return

    num_cols = len(headers) if headers else (len(rows[0]) if rows else 0)
    num_rows = len(rows) + (1 if headers else 0)
    if num_cols == 0 or num_rows == 0:
        return

    tbl_shape = slide.shapes.add_table(
        num_rows, num_cols,
        Inches(0.4), Inches(1.3), Inches(12.5), Inches(5.9),
    )
    tbl = tbl_shape.table

    if headers:
        for ci, h in enumerate(headers[:num_cols]):
            _table_cell(tbl.cell(0, ci), h,
                        size=14, bold=True,
                        text_color=Colors.WHITE, bg_color=Colors.NAVY_DARK,
                        align=PP_ALIGN.CENTER)

    row_offset = 1 if headers else 0
    for ri, row_data in enumerate(rows):
        bg = Colors.GRAY_LIGHT if ri % 2 == 0 else Colors.WHITE
        for ci, val in enumerate(row_data[:num_cols]):
            _table_cell(tbl.cell(ri + row_offset, ci), val,
                        size=13, text_color=Colors.TEXT_DARK, bg_color=bg)


def render_chart(slide, data: dict) -> None:
    """チャートスライド: 棒グラフ・折れ線グラフ対応。"""
    title = data.get("action_title") or data.get("title", "")
    _action_title_bar(slide, title)

    type_map = {
        "bar":            XL_CHART_TYPE.COLUMN_CLUSTERED,
        "bar_horizontal": XL_CHART_TYPE.BAR_CLUSTERED,
        "line":           XL_CHART_TYPE.LINE,
        "pie":            XL_CHART_TYPE.PIE,
    }
    chart_type = type_map.get(data.get("chart_type", "bar").lower(),
                               XL_CHART_TYPE.COLUMN_CLUSTERED)

    cd = ChartData()
    cd.categories = data.get("categories", [])
    for s in data.get("series", []):
        cd.add_series(s.get("name", ""), tuple(s.get("values", [])))

    chart_shape = slide.shapes.add_chart(
        chart_type,
        Inches(0.5), Inches(1.3), Inches(12.3), Inches(5.9), cd,
    )
    chart = chart_shape.chart

    # 系列カラーをネイビーパレットに変更
    for i, series in enumerate(chart.plots[0].series):
        color = CHART_PALETTE[i % len(CHART_PALETTE)]
        try:
            series.format.fill.solid()
            series.format.fill.fore_color.rgb = color
        except Exception:
            pass


def render_image(slide, data: dict) -> None:
    """画像スライド: アクションタイトルバー＋中央フィット画像＋任意キャプション。"""
    title = data.get("action_title") or data.get("title", "")
    _action_title_bar(slide, title)

    image_data = data.get("image")
    if not image_data:
        return
    src = image_data if isinstance(image_data, str) else image_data.get("src", "")
    if not src:
        return

    caption = data.get("caption", "")
    # キャプションがある場合は画像エリアを少し上に詰める
    img_top = Inches(1.3)
    img_h   = Inches(5.4) if caption else Inches(5.9)

    _add_image_fit(slide, src, Inches(0.4), img_top, Inches(12.5), img_h)

    if caption:
        _textbox(slide, caption,
                 Inches(0.4), Inches(6.95), Inches(12.5), Inches(0.4),
                 font=Fonts.BODY, size=11,
                 color=Colors.TEXT_GRAY, align=PP_ALIGN.CENTER)


def render_closing(slide, data: dict) -> None:
    """クロージングスライド: タイトルスライドに準じたレイアウト。"""
    _rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, Colors.NAVY_DARK)
    _rect(slide, Inches(0.8), Inches(3.6), Inches(4.0), Inches(0.05), Colors.BLUE_LIGHT)

    _textbox(slide, data.get("title", "ご清聴ありがとうございました"),
             Inches(0.8), Inches(2.2), Inches(11.0), Inches(1.3),
             font=Fonts.HEADING, size=32, bold=True,
             color=Colors.WHITE, align=PP_ALIGN.LEFT)

    contact = data.get("contact", "")
    if contact:
        _textbox(slide, contact,
                 Inches(0.8), Inches(3.8), Inches(10.5), Inches(0.7),
                 font=Fonts.BODY, size=16,
                 color=Colors.BLUE_LIGHT, align=PP_ALIGN.LEFT)

    note = data.get("note", "")
    if note:
        _textbox(slide, note,
                 Inches(0.8), Inches(6.5), Inches(10.5), Inches(0.6),
                 font=Fonts.BODY, size=11,
                 color=Colors.TEXT_GRAY, align=PP_ALIGN.LEFT)


# ---------------------------------------------------------------------------
# Dispatch table
# ---------------------------------------------------------------------------

RENDERERS = {
    "title":       render_title,
    "section":     render_section,
    "content":     render_content,
    "two-column":  render_two_column,
    "table":       render_table,
    "chart":       render_chart,
    "image":       render_image,
    "closing":     render_closing,
}


# ---------------------------------------------------------------------------
# Main builder
# ---------------------------------------------------------------------------

def build_presentation(slides_data: dict) -> Presentation:
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H
    blank_layout = prs.slide_layouts[6]  # blank

    for slide_data in slides_data.get("slides", []):
        layout   = slide_data.get("layout", "content").lower()
        slide    = prs.slides.add_slide(blank_layout)
        renderer = RENDERERS.get(layout)
        if renderer:
            renderer(slide, slide_data)
        else:
            print(f"WARNING: 未知のレイアウト '{layout}'。content で代替します。", file=sys.stderr)
            render_content(slide, slide_data)

    return prs


def main():
    parser = argparse.ArgumentParser(
        description="デザインシステム対応 PowerPoint ファイルを生成する",
    )
    parser.add_argument("--out",    required=True, help="出力 .pptx ファイルパス")
    parser.add_argument("--slides", help="スライド JSON ファイルパス（省略時は stdin）")
    parser.add_argument("--theme",  default=DEFAULT_THEME,
                        help=f"テーマ名（themes/<名前>.json）または JSON パス（既定: {DEFAULT_THEME}）")
    args = parser.parse_args()

    theme = load_theme(args.theme)
    if theme:
        apply_theme(theme)

    if args.slides:
        with open(args.slides, encoding="utf-8") as f:
            slides_data = json.load(f)
    else:
        slides_data = json.load(sys.stdin)

    prs = build_presentation(slides_data)
    Path(args.out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.out)

    count = len(slides_data.get("slides", []))
    print(f"生成完了: {args.out} ({count} スライド)")


if __name__ == "__main__":
    main()
